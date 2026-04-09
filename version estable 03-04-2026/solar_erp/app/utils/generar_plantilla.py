"""
Genera la plantilla oficial de importación de materiales para GeoInventario.
Uso:  from app.utils.generar_plantilla import generar_plantilla_excel
      generar_plantilla_excel("/ruta/Plantilla_GeoInventario.xlsx")
"""
from pathlib import Path


def generar_plantilla_excel(ruta_destino: str = None) -> str:
    """
    Crea el archivo Excel plantilla con formato profesional.
    Retorna la ruta del archivo generado.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.page import PageMargins
    except ImportError:
        raise RuntimeError("openpyxl no está instalado. Ejecuta: pip install openpyxl")

    if ruta_destino is None:
        from app.core.config import BASE_DIR
        ruta_destino = str(BASE_DIR.parent / "Plantilla_GeoInventario.xlsx")

    # ── Colores ──────────────────────────────────────────────────────────────
    AZUL_OSC = "1E3A5F"; AZUL_MED = "2980B9"; VERDE = "27AE60"
    VERDE_CLA = "F0FFF4"; GRIS_CLA = "F5F6FA"; BLANCO = "FFFFFF"

    CATEGORIAS = ["Paneles Solares","Inversores","Baterías","Estructuras",
                  "Cableado","Protecciones","Herramientas","Consumibles","Otros"]
    UNIDADES   = ["unidad","metro","rollo","caja","par","kg","litro",
                  "paquete","juego","bobina"]

    def _side(c="DDE2E8", s="thin"): return Side(style=s, color=c)
    def _border(c="DDE2E8"): s=_side(c); return Border(left=s,right=s,top=s,bottom=s)
    def _fill(h): return PatternFill("solid", fgColor=h)
    def _font(bold=False, color="2C3E50", size=10, italic=False):
        return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)
    def _align(h="left", wrap=False):
        return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Importar Materiales"

    # ── Cabecera del documento ────────────────────────────────────────────────
    ws.merge_cells("A1:K1")
    ws["A1"] = "🌍  GEOINVENTARIO — Plantilla de Importación de Materiales"
    ws["A1"].font = Font(name="Arial", bold=True, size=15, color=BLANCO)
    ws["A1"].fill = _fill(AZUL_OSC)
    ws["A1"].alignment = _align("center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:K2")
    ws["A2"] = "Ing. Mateo Salazar Zapata  |  Área de Ingeniería Solar"
    ws["A2"].font = Font(name="Arial", italic=True, size=10, color=BLANCO)
    ws["A2"].fill = _fill(AZUL_MED)
    ws["A2"].alignment = _align("center")
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:K3")
    ws["A3"].fill = _fill(GRIS_CLA)
    ws.row_dimensions[3].height = 8

    ws.merge_cells("A4:K4")
    ws["A4"] = (
        "📋  INSTRUCCIONES: Completa los campos (*) obligatorios. "
        "Usa las listas desplegables en Categoría y Unidad. "
        "El campo 'precio_total' se calcula automáticamente. "
        "No modifiques la fila 7 de encabezados."
    )
    ws["A4"].font = Font(name="Arial", size=9, italic=True, color="5D4037")
    ws["A4"].fill = _fill("FFF8E1")
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 34

    ws.merge_cells("A5:K5")
    ws["A5"].fill = _fill(GRIS_CLA)
    ws.row_dimensions[5].height = 6

    # ── Definición de columnas ────────────────────────────────────────────────
    COLS = [
        ("A", "nombre_material *",   "Nombre completo del material o componente",                    18, True),
        ("B", "categoria *",          "Categoría (lista desplegable)",                                18, True),
        ("C", "unidad *",             "Unidad de medida (lista desplegable)",                         12, True),
        ("D", "cantidad_actual",      "Cantidad en existencia actualmente",                           14, False),
        ("E", "stock_minimo",         "Cantidad mínima antes de alerta",                              14, False),
        ("F", "precio_unitario",      "Precio de compra por unidad en COP ($)",                       16, False),
        ("G", "precio_total",         "= cantidad_actual × precio_unitario  [AUTOMÁTICO]",            18, False),
        ("H", "ubicacion",            "Bodega, estante o zona de almacenamiento",                     18, False),
        ("I", "proveedor_nit",        "NIT del proveedor (debe existir en el sistema)",               15, False),
        ("J", "observaciones",        "Notas adicionales o descripción",                              24, False),
    ]

    # Fila 6: subtítulos descriptivos
    ws.row_dimensions[6].height = 28
    for col_letter, _hdr, desc, _w, _ob in COLS:
        col_idx = ord(col_letter) - 64
        c = ws.cell(row=6, column=col_idx, value=desc)
        c.font = Font(name="Arial", size=8, italic=True, color="546E7A")
        c.fill = _fill("ECEFF1")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border("CFD8DC")

    # Fila 7: encabezados principales
    ws.row_dimensions[7].height = 36
    for col_letter, hdr, _desc, width, obligatorio in COLS:
        col_idx = ord(col_letter) - 64
        ws.column_dimensions[col_letter].width = width
        c = ws.cell(row=7, column=col_idx, value=hdr)
        c.font = Font(name="Arial", bold=True, size=10, color=BLANCO)
        c.fill = _fill(AZUL_OSC if obligatorio else AZUL_MED)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(
            left  =Side(style="medium", color=BLANCO),
            right =Side(style="medium", color=BLANCO),
            bottom=Side(style="medium", color=BLANCO),
        )

    ws["G7"].value = "precio_total\n[automático 🔒]"
    ws["G7"].fill = _fill("1A5276")

    # ── Datos de ejemplo ──────────────────────────────────────────────────────
    EJEMPLOS = [
        ("Panel Solar 550W Monocristalino", "Paneles Solares", "unidad",  10, 2,  320000),
        ("Inversor Huawei SUN2000 5KTL",    "Inversores",     "unidad",   3,  1, 2800000),
        ("Cable Solar 4mm² Rojo",           "Cableado",       "metro",   200, 50,   3500),
        ("Estructura Ballast 12 paneles",   "Estructuras",    "juego",    5,  1,  185000),
        ("Breaker DC 15A",                  "Protecciones",   "unidad",  20,  5,   28000),
    ]
    for i, (nombre, cat, uni, cant, mini, precio) in enumerate(EJEMPLOS):
        r = 8 + i
        ws.row_dimensions[r].height = 22
        bg = BLANCO if i % 2 == 0 else GRIS_CLA
        data = [nombre, cat, uni, cant, mini, precio]
        for ci, v in enumerate(data, 1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = Font(name="Arial", size=10, italic=True, color="2C3E50")
            c.fill = _fill(bg)
            c.border = _border()
            c.alignment = _align("left" if ci == 1 else "center")
        # precio_total fórmula
        pt = ws.cell(row=r, column=7, value=f"=D{r}*F{r}")
        pt.font = Font(name="Arial", bold=True, size=10, color=VERDE)
        pt.fill = _fill(VERDE_CLA)
        pt.number_format = '"$"#,##0.00'
        pt.border = _border(VERDE)
        pt.alignment = _align("right")
        # formatos numéricos
        for ci in (4, 5): ws.cell(row=r, column=ci).number_format = "#,##0.00"
        ws.cell(row=r, column=6).number_format = '"$"#,##0.00'
        for ci in (8, 9, 10):
            ws.cell(row=r, column=ci).fill = _fill(bg); ws.cell(row=r, column=ci).border = _border()

    # ── Filas vacías (13-62) ──────────────────────────────────────────────────
    for r in range(13, 63):
        ws.row_dimensions[r].height = 22
        bg = BLANCO if r % 2 == 0 else GRIS_CLA
        for ci in range(1, 11):
            c = ws.cell(row=r, column=ci)
            c.fill = _fill(bg); c.border = _border()
            c.font = Font(name="Arial", size=10)
            c.alignment = _align("left" if ci in (1, 8, 9, 10) else "center")
        pt = ws.cell(row=r, column=7, value=f'=IF(OR(D{r}="",D{r}=0),"",D{r}*F{r})')
        pt.font = Font(name="Arial", bold=True, size=10, color=VERDE)
        pt.fill = _fill(VERDE_CLA if r % 2 == 0 else "E8F8F5")
        pt.number_format = '"$"#,##0.00'
        pt.border = _border(VERDE)
        pt.alignment = _align("right")
        for ci in (4, 5): ws.cell(row=r, column=ci).number_format = "#,##0.00"
        ws.cell(row=r, column=6).number_format = '"$"#,##0.00'

    # ── Fila de totales (63) ──────────────────────────────────────────────────
    rt = 63
    ws.row_dimensions[rt].height = 28
    ws.merge_cells(f"A{rt}:C{rt}")
    c = ws.cell(row=rt, column=1, value="TOTALES INVENTARIO")
    c.font = Font(name="Arial", bold=True, size=11, color=BLANCO)
    c.fill = _fill(AZUL_OSC); c.alignment = _align("center")
    for ci in (4, 5):
        c = ws.cell(row=rt, column=ci, value=f"=SUM({get_column_letter(ci)}8:{get_column_letter(ci)}{rt-1})")
        c.font = Font(name="Arial", bold=True, color=BLANCO)
        c.fill = _fill(AZUL_MED); c.number_format = "#,##0.00"; c.alignment = _align("center")
    ws.cell(row=rt, column=6).fill = _fill(AZUL_OSC)
    ct = ws.cell(row=rt, column=7, value=f"=SUM(G8:G{rt-1})")
    ct.font = Font(name="Arial", bold=True, size=12, color=BLANCO)
    ct.fill = _fill(VERDE); ct.number_format = '"$"#,##0.00'; ct.alignment = _align("right")
    for ci in (8, 9, 10): ws.cell(row=rt, column=ci).fill = _fill(AZUL_OSC)

    # ── Validaciones de datos ─────────────────────────────────────────────────
    dv_cat = DataValidation(type="list",
        formula1='"' + ",".join(CATEGORIAS) + '"',
        allow_blank=True, showDropDown=False, showErrorMessage=True,
        errorTitle="Categoría inválida",
        error="Selecciona una categoría de la lista desplegable.")
    ws.add_data_validation(dv_cat)
    dv_cat.sqref = "B8:B62"

    dv_uni = DataValidation(type="list",
        formula1='"' + ",".join(UNIDADES) + '"',
        allow_blank=True, showDropDown=False, showErrorMessage=False)
    ws.add_data_validation(dv_uni)
    dv_uni.sqref = "C8:C62"

    dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual",
        formula1="0", allow_blank=True, showErrorMessage=True,
        errorTitle="Valor inválido", error="Debe ser un número ≥ 0.")
    ws.add_data_validation(dv_num)
    dv_num.sqref = "D8:F62"

    # ── Freeze / autofilter ───────────────────────────────────────────────────
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:J{rt-1}"

    # ── Hoja de Referencia ────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Referencia")
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "Valores de referencia para la importación — GeoInventario"
    ws2["A1"].font = Font(name="Arial", bold=True, size=12, color=BLANCO)
    ws2["A1"].fill = _fill(AZUL_OSC)
    ws2["A1"].alignment = _align("center")
    ws2.row_dimensions[1].height = 26

    for ci, (title, w) in enumerate([(( "Categorías válidas", 22)), (("Unidades válidas", 18))], 1):
        c = ws2.cell(row=2, column=ci*2-1, value=title)
        c.font = Font(name="Arial", bold=True, color=BLANCO)
        c.fill = _fill(AZUL_MED); c.alignment = _align("center"); c.border = _border()
        ws2.column_dimensions[get_column_letter(ci*2-1)].width = w

    for i, cat in enumerate(CATEGORIAS, 3):
        c = ws2.cell(row=i, column=1, value=cat)
        c.fill = _fill(BLANCO if i%2==0 else GRIS_CLA); c.border = _border()
    for i, uni in enumerate(UNIDADES, 3):
        c = ws2.cell(row=i, column=3, value=uni)
        c.fill = _fill(BLANCO if i%2==0 else GRIS_CLA); c.border = _border()

    notas = [
        "Notas importantes:",
        "• El campo precio_total es AUTOMÁTICO (cantidad × precio_unitario).",
        "• El sistema internamente calcula precio_total al importar.",
        "• proveedor_nit debe coincidir con un NIT registrado en el sistema.",
        "• Las categorías no reconocidas se importan tal como están.",
        "• Deja precio_unitario en 0 si no tienes el dato (puedes editarlo después).",
    ]
    for j, nota in enumerate(notas, 15):
        c = ws2.cell(row=j, column=1, value=nota)
        c.font = Font(name="Arial", size=9,
                      bold=(j == 15), italic=(j > 15),
                      color=("E74C3C" if j == 15 else "546E7A"))

    ws2.column_dimensions["B"].width = 3
    ws2.column_dimensions["D"].width = 3

    # ── Configuración de página ───────────────────────────────────────────────
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)
    ws.print_title_rows = "7:7"

    Path(ruta_destino).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_destino)
    return ruta_destino

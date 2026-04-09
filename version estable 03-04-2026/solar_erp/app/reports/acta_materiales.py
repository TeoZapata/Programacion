from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from datetime import date
import os
from pathlib import Path


def generar_acta_materiales(movimiento, proyecto, cliente, detalles, responsable, output_dir=None):
    """
    Genera un documento Word con el acta de entrega de materiales.
    Retorna la ruta del archivo generado.
    """
    doc = Document()

    # Márgenes
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ---- ENCABEZADO ----
    titulo = doc.add_heading("", level=0)
    run = titulo.add_run("🌍  GEOINVENTARIO")
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    run.font.size = Pt(20)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitulo.add_run("ACTA DE ENTREGA DE MATERIALES")
    sub_run.font.bold = True
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(0x2E, 0xCC, 0x71)

    doc.add_paragraph()

    # ---- INFO DEL PROYECTO ----
    tabla_info = doc.add_table(rows=5, cols=2)
    tabla_info.style = "Table Grid"

    datos = [
        ("Proyecto:", proyecto.nombre_proyecto if proyecto else "—"),
        ("Cliente:", cliente.nombre if cliente else "—"),
        ("Dirección:", proyecto.direccion if proyecto else "—"),
        ("Fecha:", str(movimiento.fecha or date.today())),
        ("Responsable:", responsable or movimiento.responsable or "—"),
    ]

    for i, (etiqueta, valor) in enumerate(datos):
        fila = tabla_info.rows[i]
        celda_etiq = fila.cells[0]
        celda_val = fila.cells[1]
        celda_etiq.width = Cm(4)
        celda_val.width = Cm(11)
        p_etiq = celda_etiq.paragraphs[0]
        r = p_etiq.add_run(etiqueta)
        r.bold = True
        r.font.size = Pt(10)
        p_val = celda_val.paragraphs[0]
        p_val.add_run(str(valor)).font.size = Pt(10)

    doc.add_paragraph()

    # ---- TABLA DE MATERIALES ----
    heading_mat = doc.add_heading("Materiales Entregados", level=2)
    heading_mat.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    cols_header = ["#", "Material", "Unidad", "Cantidad"]
    tabla_mat = doc.add_table(rows=1, cols=len(cols_header))
    tabla_mat.style = "Table Grid"

    # Cabecera
    hdr_row = tabla_mat.rows[0]
    for i, col in enumerate(cols_header):
        cell = hdr_row.cells[i]
        p = cell.paragraphs[0]
        run = p.add_run(col)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        from docx.oxml.ns import qn
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "1E3A5F")
        tcPr.append(shd)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Filas de materiales
    total_items = len(detalles)
    for i, detalle in enumerate(detalles):
        fila = tabla_mat.add_row()
        mat = detalle.material
        valores = [
            str(i + 1),
            mat.nombre_material if mat else "—",
            mat.unidad if mat else "—",
            str(detalle.cantidad),
        ]
        for j, val in enumerate(valores):
            cell = fila.cells[j]
            p = cell.paragraphs[0]
            p.add_run(val).font.size = Pt(10)
            if j == 0 or j == 2 or j == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Color alternado
            if i % 2 == 1:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "EBF5FB")
                tcPr.append(shd)

    doc.add_paragraph()

    # ---- OBSERVACIONES ----
    if movimiento.observaciones:
        doc.add_heading("Observaciones", level=3)
        doc.add_paragraph(movimiento.observaciones)

    doc.add_paragraph()
    doc.add_paragraph()

    # ---- FIRMAS ----
    tabla_firmas = doc.add_table(rows=3, cols=2)
    firmas = [
        ("Entregado por:", "Recibido por:"),
        ("", ""),
        (responsable or "___________________", "___________________"),
    ]
    for i, (izq, der) in enumerate(firmas):
        fila = tabla_firmas.rows[i]
        p_izq = fila.cells[0].paragraphs[0]
        p_der = fila.cells[1].paragraphs[0]
        if i == 0:
            p_izq.add_run(izq).bold = True
            p_der.add_run(der).bold = True
        else:
            p_izq.add_run(izq)
            p_der.add_run(der)
        p_izq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_der.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- PIE ----
    doc.add_paragraph()
    pie = doc.add_paragraph(f"Generado por GeoInventario — Ing. Mateo Salazar Zapata  •  {date.today().strftime('%d/%m/%Y')}")
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    pie.runs[0].font.size = Pt(8)

    # ---- GUARDAR ----
    if output_dir is None:
        from app.core.config import BASE_DIR
        output_dir = BASE_DIR / "documentos"
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    nombre_archivo = f"Acta_Materiales_{proyecto.nombre_proyecto if proyecto else 'Sin_Proyecto'}_{movimiento.id}.docx"
    nombre_archivo = nombre_archivo.replace(" ", "_").replace("/", "-")
    ruta = Path(output_dir) / nombre_archivo
    doc.save(str(ruta))
    return str(ruta)

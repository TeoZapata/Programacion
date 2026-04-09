"""
Exportación de inventario en formato SIIGO (.xlsx)
Genera el archivo con la plantilla oficial de importación de productos SIIGO.
"""
from pathlib import Path
from datetime import date


# Mapeo de unidades GeoInventario → código DIAN SIIGO
UNIDAD_A_DIAN = {
    "unidad":   "94",   # Unidad
    "metro":    "MTR",  # Metro
    "rollo":    "94",
    "caja":     "BX",   # Caja
    "par":      "PR",   # Par
    "kg":       "KGM",  # Kilogramo
    "litro":    "LT",   # Litro
    "paquete":  "PK",   # Paquete
    "juego":    "SET",  # Juego/Set
    "bobina":   "RL",   # Rollo
    "":         "94",
}

# Mapeo de categorías GeoInventario → categoría SIIGO
CATEGORIA_SIIGO = "1 Productos"


def exportar_siigo(materiales, output_dir=None) -> str:
    """
    Exporta la lista de materiales al formato de importación SIIGO.
    Retorna la ruta del archivo generado.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl no está instalado.")

    if output_dir is None:
        from app.core.config import BASE_DIR
        output_dir = BASE_DIR / "documentos"
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    fecha_str = date.today().strftime("%Y%m%d")
    ruta = Path(output_dir) / f"Exportacion_SIIGO_{fecha_str}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"

    # ── Colores ────────────────────────────────────────────────────────────
    AZUL_HDR  = "1E3A5F"
    VERDE_HDR = "27AE60"
    GRIS_CLA  = "F5F6FA"
    BLANCO    = "FFFFFF"

    def _side(c="DDE2E8"):
        return Border(
            left  =Side(style="thin", color=c),
            right =Side(style="thin", color=c),
            top   =Side(style="thin", color=c),
            bottom=Side(style="thin", color=c),
        )

    # ── Encabezados — exactamente los de la plantilla SIIGO ───────────────
    HEADERS = [
        "Tipo de Producto (obligatorio)",
        "Categoría de Inventarios / Servicios (obligatorio)",
        "Código del Producto (obligatorio)",
        "Nombre del Producto / Servicio (obligatorio)",
        "¿Inventariable? (obligatorio)",
        "Visible en facturas de venta",
        "Stock mínimo",
        "Código Unidad de medida DIAN",
        "Unidad de Medida Impresión Factura",
        "Referencia de Fábrica",
        "Código de Barras",
        "Descripción Larga",
        "Código Impuesto Retención",
        "Código Impuesto Cargo",
        "Código Impuesto Cargo Dos",
        "Cant. de mililitros",
        "Tarifa",
        "Valor Impuesto Cargo Dos",
        "¿Incluye IVA en Precio de Venta?",
        "Precio de venta 1",
        "Precio de venta 2",
        "Precio de venta 3",
        "Precio de venta 4",
        "Precio de venta 5",
        "Precio de venta 6",
        "Precio de venta 7",
        "Precio de venta 8",
        "Precio de venta 9",
        "Precio de venta 10",
        "Precio de venta 11",
        "Precio de venta 12",
        "Código Arancelario",
        "Marca",
        "Modelo",
    ]

    # Fila 1 — encabezados con estilo
    ws.row_dimensions[1].height = 38
    for col_idx, header in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.font      = Font(name="Arial", bold=True, size=9, color=BLANCO)
        c.fill      = PatternFill("solid", fgColor=AZUL_HDR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = _side("2980B9")

    # Anchos de columna aproximados
    COL_WIDTHS = [18, 20, 14, 35, 14, 14, 10, 12, 14, 14, 14, 20,
                  12, 12, 12, 10, 10, 10, 14,
                  12, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10,
                  14, 14, 14]
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Filas de datos ─────────────────────────────────────────────────────
    for row_idx, mat in enumerate(materiales, start=2):
        bg = BLANCO if row_idx % 2 == 0 else GRIS_CLA
        ws.row_dimensions[row_idx].height = 18

        unidad_norm = (mat.unidad or "").lower().strip()
        cod_dian    = UNIDAD_A_DIAN.get(unidad_norm, "94")
        codigo      = str(mat.id)   # ID de la BD — identificador único del producto
        precio_v    = mat.precio_unitario or 0

        row_data = [
            "P-Producto",           # A: Tipo
            CATEGORIA_SIIGO,        # B: Categoría
            codigo,                 # C: Código
            mat.nombre_material,    # D: Nombre
            "SI",                   # E: ¿Inventariable?
            "SI",                   # F: Visible en facturas
            mat.stock_minimo or 0,  # G: Stock mínimo
            cod_dian,               # H: Código DIAN
            mat.unidad or "unidad", # I: Unidad impresión
            None,                   # J: Referencia fábrica
            None,                   # K: Código barras
            mat.ubicacion or None,  # L: Descripción larga
            None,                   # M: Imp. Retención
            None,                   # N: Imp. Cargo
            None,                   # O: Imp. Cargo Dos
            None,                   # P: Mililitros
            None,                   # Q: Tarifa
            None,                   # R: Valor Imp. Cargo Dos
            "NO",                   # S: ¿Incluye IVA?
            precio_v if precio_v > 0 else None,  # T: Precio venta 1
        ] + [None] * 14             # U-AH: Precios 2-12, arancelario, marca, modelo

        for col_idx, val in enumerate(row_data, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font      = Font(name="Arial", size=9)
            c.fill      = PatternFill("solid", fgColor=bg)
            c.border    = _side()
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if col_idx in (4, 12) else "center")

        # Formato numérico para precio
        if precio_v > 0:
            ws.cell(row=row_idx, column=20).number_format = '"$"#,##0.00'
        ws.cell(row=row_idx, column=7).number_format = "#,##0.##"

    # ── Fila de totales info ───────────────────────────────────────────────
    r_tot = len(materiales) + 2
    ws.row_dimensions[r_tot].height = 22
    ws.merge_cells(f"A{r_tot}:C{r_tot}")
    c = ws.cell(row=r_tot, column=1, value=f"Total productos exportados: {len(materiales)}")
    c.font      = Font(name="Arial", bold=True, size=10, color=BLANCO)
    c.fill      = PatternFill("solid", fgColor=AZUL_HDR)
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"D{r_tot}:F{r_tot}")
    c2 = ws.cell(row=r_tot, column=4,
                 value=f"Generado por GeoInventario — {date.today().strftime('%d/%m/%Y')}")
    c2.font      = Font(name="Arial", italic=True, size=9, color=BLANCO)
    c2.fill      = PatternFill("solid", fgColor="2980B9")
    c2.alignment = Alignment(horizontal="center", vertical="center")

    # ── Hoja Listas (requerida por SIIGO) ─────────────────────────────────
    ws2 = wb.create_sheet("Listas")
    listas = [
        ["P-Producto", "S-Servicio"],
        ["1 Productos", "2 Servicios"],
        ["SI", "NO"],
    ]
    for r, fila in enumerate(listas, start=1):
        for c_idx, val in enumerate(fila, start=1):
            ws2.cell(row=r, column=c_idx, value=val)

    # ── Freeze y autofilter ────────────────────────────────────────────────
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    wb.save(str(ruta))
    return str(ruta)

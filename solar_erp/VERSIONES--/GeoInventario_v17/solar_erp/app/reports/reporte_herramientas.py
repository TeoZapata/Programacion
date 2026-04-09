"""
Generador de reportes Word para herramientas — individual y general.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import os
from pathlib import Path


# ── Paleta ────────────────────────────────────────────────────────────────────
AZUL   = RGBColor(0x1E, 0x3A, 0x5F)
VERDE  = RGBColor(0x27, 0xAE, 0x60)
GRIS   = RGBColor(0x71, 0x8C, 0x96)
NEGRO  = RGBColor(0x2C, 0x3E, 0x50)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

AZUL_HEX  = "1E3A5F"
VERDE_HEX = "27AE60"
GRIS_HEX  = "ECF0F1"


def _set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_text(cell, texto, bold=False, size=10, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(str(texto))
    run.font.size  = Pt(size)
    run.font.bold  = bold
    if color:
        run.font.color.rgb = color


def _encabezado(doc, titulo_reporte):
    section = doc.sections[0]
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

    t = doc.add_heading("", level=0)
    r = t.add_run("🌍  GEOINVENTARIO")
    r.font.color.rgb = AZUL
    r.font.size = Pt(18)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(titulo_reporte)
    sr.font.bold = True
    sr.font.size = Pt(13)
    sr.font.color.rgb = VERDE

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  "
                      "Área de Ingeniería — Ing. Mateo Salazar Zapata")
    mr.font.size = Pt(9)
    mr.font.color.rgb = GRIS

    doc.add_paragraph()


def _tabla_herramientas(doc, herramientas, titulo_seccion=""):
    if titulo_seccion:
        h = doc.add_heading(titulo_seccion, level=2)
        h.runs[0].font.color.rgb = AZUL

    cols = ["#", "Código", "Nombre", "Categoría", "Marca/Modelo", "N° Serie", "Estado", "Ubicación"]
    widths = [Cm(0.8), Cm(2.2), Cm(4.5), Cm(2.5), Cm(3), Cm(2.5), Cm(2.3), Cm(2.8)]

    tabla = doc.add_table(rows=1, cols=len(cols))
    tabla.style = "Table Grid"

    # Encabezados
    for i, (col, w) in enumerate(zip(cols, widths)):
        cell = tabla.cell(0, i)
        cell.width = w
        _set_cell_bg(cell, AZUL_HEX)
        _cell_text(cell, col, bold=True, size=9, color=BLANCO, align=WD_ALIGN_PARAGRAPH.CENTER)

    ESTADO_EMOJI = {
        "disponible":   "✅ Disponible",
        "asignada":     "🔧 Asignada",
        "mantenimiento":"🔄 Mantenim.",
        "baja":         "❌ Baja",
    }
    for idx, h in enumerate(herramientas, 1):
        row = tabla.add_row()
        bg  = GRIS_HEX if idx % 2 == 0 else "FFFFFF"
        vals = [
            str(idx),
            h.codigo or "—",
            h.nombre,
            h.categoria or "—",
            f"{h.marca or ''} {h.modelo or ''}".strip() or "—",
            h.numero_serie or "—",
            ESTADO_EMOJI.get(h.estado, h.estado),
            h.ubicacion or "—",
        ]
        for ci, v in enumerate(vals):
            c = row.cells[ci]
            _set_cell_bg(c, bg)
            _cell_text(c, v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER if ci == 0 else WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph()
    return len(herramientas)


def _tabla_asignaciones(doc, asignaciones, titulo_seccion=""):
    if titulo_seccion:
        h = doc.add_heading(titulo_seccion, level=2)
        h.runs[0].font.color.rgb = AZUL

    if not asignaciones:
        p = doc.add_paragraph("Sin asignaciones activas.")
        p.runs[0].font.italic = True
        return

    cols = ["#", "Herramienta", "Código", "Categoría", "Trabajador", "Fecha Asignación", "Observaciones"]
    widths = [Cm(0.7), Cm(4.5), Cm(2.2), Cm(2.5), Cm(3.5), Cm(2.8), Cm(4)]

    tabla = doc.add_table(rows=1, cols=len(cols))
    tabla.style = "Table Grid"

    for i, (col, w) in enumerate(zip(cols, widths)):
        cell = tabla.cell(0, i)
        cell.width = w
        _set_cell_bg(cell, VERDE_HEX)
        _cell_text(cell, col, bold=True, size=9, color=BLANCO, align=WD_ALIGN_PARAGRAPH.CENTER)

    for idx, a in enumerate(asignaciones, 1):
        row = tabla.add_row()
        bg  = GRIS_HEX if idx % 2 == 0 else "FFFFFF"
        nombre_herr   = a.herramienta.nombre if a.herramienta else "—"
        codigo_herr   = a.herramienta.codigo if a.herramienta else "—"
        cat_herr      = a.herramienta.categoria if a.herramienta else "—"
        nombre_trab   = a.trabajador.nombre if a.trabajador else "—"
        fecha         = a.fecha_asignacion.strftime("%d/%m/%Y %H:%M") if a.fecha_asignacion else "—"
        vals = [str(idx), nombre_herr, codigo_herr or "—", cat_herr or "—",
                nombre_trab, fecha, a.observaciones or ""]
        for ci, v in enumerate(vals):
            c = row.cells[ci]
            _set_cell_bg(c, bg)
            _cell_text(c, v, size=9)

    doc.add_paragraph()


def _pie(doc, n_herr=None, n_asig=None):
    doc.add_paragraph()
    linea = doc.add_paragraph("─" * 90)
    linea.runs[0].font.color.rgb = GRIS
    linea.runs[0].font.size = Pt(8)

    resumen_parts = []
    if n_herr is not None:
        resumen_parts.append(f"Total herramientas: {n_herr}")
    if n_asig is not None:
        resumen_parts.append(f"Asignaciones activas: {n_asig}")

    if resumen_parts:
        p = doc.add_paragraph("  |  ".join(resumen_parts))
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = AZUL

    firma = doc.add_paragraph(
        "Firma responsable: ____________________________     "
        "Cargo: ____________________________     "
        "Fecha: ________________"
    )
    firma.runs[0].font.size = Pt(9)
    firma.runs[0].font.color.rgb = GRIS


# ── Reporte General ───────────────────────────────────────────────────────────
def generar_reporte_general(herramientas, asignaciones_activas, output_dir=None):
    if output_dir is None:
        output_dir = Path.home() / "GeoInventario" / "reportes"
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    doc  = Document()
    _encabezado(doc, "REPORTE GENERAL DE HERRAMIENTAS")

    # Resumen ejecutivo
    disponibles   = sum(1 for h in herramientas if h.estado == "disponible")
    asignadas     = sum(1 for h in herramientas if h.estado == "asignada")
    mantenimiento = sum(1 for h in herramientas if h.estado == "mantenimiento")
    bajas         = sum(1 for h in herramientas if h.estado == "baja")

    res = doc.add_heading("Resumen Ejecutivo", level=2)
    res.runs[0].font.color.rgb = AZUL

    tabla_res = doc.add_table(rows=2, cols=4)
    tabla_res.style = "Table Grid"
    headers_r = ["Disponibles", "Asignadas", "En Mantenimiento", "Bajas"]
    colores_r = [VERDE_HEX, "2980B9", "F39C12", "E74C3C"]
    valores_r = [disponibles, asignadas, mantenimiento, bajas]

    for i, (h, c, v) in enumerate(zip(headers_r, colores_r, valores_r)):
        cell_h = tabla_res.cell(0, i)
        _set_cell_bg(cell_h, c)
        _cell_text(cell_h, h, bold=True, size=10, color=BLANCO, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_v = tabla_res.cell(1, i)
        _cell_text(cell_v, str(v), bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()

    n = _tabla_herramientas(doc, herramientas, "Inventario Completo de Herramientas")
    _tabla_asignaciones(doc, asignaciones_activas, "Asignaciones Activas")
    _pie(doc, n_herr=len(herramientas), n_asig=len(asignaciones_activas))

    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = Path(output_dir) / f"Reporte_Herramientas_General_{ts}.docx"
    doc.save(str(path))
    return str(path)


# ── Reporte Individual (por trabajador) ───────────────────────────────────────
def generar_reporte_trabajador(trabajador, asignaciones, output_dir=None):
    if output_dir is None:
        output_dir = Path.home() / "GeoInventario" / "reportes"
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    doc = Document()
    nombre_trab = trabajador.nombre if trabajador else "Trabajador"
    _encabezado(doc, f"REPORTE DE HERRAMIENTAS — {nombre_trab.upper()}")

    # Info del trabajador
    h2 = doc.add_heading("Datos del Trabajador", level=2)
    h2.runs[0].font.color.rgb = AZUL

    tabla_t = doc.add_table(rows=3, cols=2)
    tabla_t.style = "Table Grid"
    datos = [
        ("Nombre completo", trabajador.nombre if trabajador else "—"),
        ("Cargo",           trabajador.cargo   if trabajador else "—"),
        ("Teléfono",        trabajador.telefono if trabajador else "—"),
    ]
    for row_i, (lbl, val) in enumerate(datos):
        c0 = tabla_t.cell(row_i, 0)
        c1 = tabla_t.cell(row_i, 1)
        _set_cell_bg(c0, AZUL_HEX)
        _cell_text(c0, lbl, bold=True, size=10, color=BLANCO)
        _cell_text(c1, val, size=10)

    doc.add_paragraph()

    herramientas_asig = [a.herramienta for a in asignaciones if a.herramienta]
    _tabla_herramientas(doc, herramientas_asig,
                        f"Herramientas Asignadas ({len(herramientas_asig)})")
    _tabla_asignaciones(doc, asignaciones, "Detalle de Asignaciones")
    _pie(doc, n_asig=len(asignaciones))

    ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = "".join(c for c in nombre_trab if c.isalnum() or c in " _-")[:30]
    path = Path(output_dir) / f"Reporte_Herramientas_{safe}_{ts}.docx"
    doc.save(str(path))
    return str(path)

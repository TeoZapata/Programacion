"""
Utilidad para importar herramientas desde plantilla Excel.
"""


def leer_plantilla_herramientas(ruta_archivo: str):
    try:
        import openpyxl
    except ImportError:
        return [], ["openpyxl no instalado. Ejecuta: pip install openpyxl"], []

    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
    except Exception as e:
        return [], [f"No se pudo abrir el archivo: {e}"], []

    hoja = None
    for nombre in wb.sheetnames:
        if "herramienta" in nombre.lower() or "importar" in nombre.lower():
            hoja = wb[nombre]
            break
    if hoja is None:
        hoja = wb.active

    # Detectar fila de encabezados buscando "nombre_herramienta"
    fila_enc = None
    for row in hoja.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if cell.value and "nombre_herramienta" in str(cell.value).lower():
                fila_enc = cell.row
                break
        if fila_enc:
            break

    if fila_enc is None:
        return [], ["No se encontró la fila de encabezados. Usa la plantilla oficial."], []

    col_map = {}
    for cell in hoja[fila_enc]:
        if cell.value:
            col_map[str(cell.value).strip().lower()] = cell.column

    def gc(key):
        for k, v in col_map.items():
            if key in k:
                return v
        return None

    col_nombre  = gc("nombre_herramienta")
    col_codigo  = gc("codigo")
    col_cat     = gc("categoria")
    col_marca   = gc("marca")
    col_modelo  = gc("modelo")
    col_serie   = gc("numero_serie") or gc("serie")
    col_estado  = gc("estado")
    col_ubic    = gc("ubicacion")
    col_obs     = gc("observaciones")

    if not col_nombre:
        return [], ["No se encontró la columna 'nombre_herramienta'. Usa la plantilla oficial."], []

    ESTADOS_VALIDOS = {"disponible", "asignada", "mantenimiento", "baja"}
    CATEGORIAS_VALIDAS = {
        "herramienta eléctrica", "herramienta manual", "medición",
        "seguridad", "instalación solar", "elevación", "general",
    }

    items_ok     = []
    errores      = []
    advertencias = []

    for row_idx in range(fila_enc + 1, hoja.max_row + 1):
        def val(col):
            if col is None:
                return None
            v = hoja.cell(row=row_idx, column=col).value
            return str(v).strip() if v is not None else None

        nombre  = val(col_nombre)
        if not nombre:
            continue

        codigo  = val(col_codigo)
        cat     = val(col_cat)
        marca   = val(col_marca)
        modelo  = val(col_modelo)
        serie   = val(col_serie)
        estado  = val(col_estado)
        ubic    = val(col_ubic)
        obs     = val(col_obs)

        fila_label = f"Fila {row_idx}"
        fila_errs  = []

        if not estado:
            estado = "disponible"
        elif estado.lower() not in ESTADOS_VALIDOS:
            advertencias.append(f"{fila_label}: Estado '{estado}' no reconocido, se usará 'disponible'.")
            estado = "disponible"

        if cat and cat.lower() not in CATEGORIAS_VALIDAS:
            advertencias.append(f"{fila_label}: Categoría '{cat}' no está en la lista estándar.")

        if fila_errs:
            errores.extend(fila_errs)
            continue

        items_ok.append({
            "codigo":       codigo or "",
            "nombre":       nombre,
            "categoria":    cat or "General",
            "marca":        marca or "",
            "modelo":       modelo or "",
            "numero_serie": serie or "",
            "estado":       estado.lower(),
            "ubicacion":    ubic or "",
            "observaciones":obs or "",
        })

    return items_ok, errores, advertencias

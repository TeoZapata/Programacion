"""
==============================================================
backend/services/report_service.py
Generación de reportes en Excel para certificados RETIE
==============================================================
"""

import io
import logging
from datetime import date, datetime, timedelta

logger = logging.getLogger(__name__)


class ReportService:
    """Servicio para generar reportes y exportaciones."""

    def exportar_excel_completo(self, certificados: list) -> bytes:
        """
        Exporta todos los certificados a Excel (.xlsx).
        Returns: bytes del archivo Excel
        """
        try:
            import openpyxl
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side
            )
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise RuntimeError("openpyxl no instalado. Ejecuta: pip install openpyxl")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Certificados RETIE"

        # ── Estilos ────────────────────────────────────────
        header_fill = PatternFill("solid", fgColor="1B4F8A")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        center = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        fill_vigente = PatternFill("solid", fgColor="D5F5E3")
        fill_vencer = PatternFill("solid", fgColor="FEF9E7")
        fill_vencido = PatternFill("solid", fgColor="FDEDEC")

        # ── Título ─────────────────────────────────────────
        ws.merge_cells("A1:J1")
        title_cell = ws["A1"]
        title_cell.value = f"REPORTE DE CERTIFICADOS RETIE — Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        title_cell.font = Font(bold=True, size=13, color="1B4F8A")
        title_cell.alignment = center
        ws.row_dimensions[1].height = 28

        # ── Encabezados ────────────────────────────────────
        headers = [
            "N° Certificado", "Producto", "Organismo Certificador",
            "Convenio", "Descripción", "Fecha Emisión",
            "Fecha Vencimiento", "Estado", "Días para Vencer", "Notas"
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

        ws.row_dimensions[2].height = 22

        # ── Datos ──────────────────────────────────────────
        for row_idx, cert in enumerate(certificados, 3):
            dias = cert.dias_para_vencer
            estado = cert.estado_certificado

            row_data = [
                cert.numero_certificado,
                cert.producto,
                cert.organismo_certificador,
                cert.convenio or "",
                cert.descripcion or "",
                cert.fecha_emision.strftime("%d/%m/%Y") if cert.fecha_emision else "",
                cert.fecha_vencimiento.strftime("%d/%m/%Y") if cert.fecha_vencimiento else "Sin fecha",
                estado.upper().replace("_", " "),
                dias if dias is not None else "N/A",
                cert.notas or "",
            ]

            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

                # Color de fila según estado
                if estado == "vigente":
                    cell.fill = fill_vigente
                elif estado == "por_vencer":
                    cell.fill = fill_vencer
                elif estado == "vencido":
                    cell.fill = fill_vencido

        # ── Anchos de columna ──────────────────────────────
        col_widths = [22, 30, 30, 20, 40, 15, 18, 14, 16, 30]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # ── Filtros automáticos ────────────────────────────
        ws.auto_filter.ref = f"A2:J{len(certificados) + 2}"

        # ── Hoja de resumen ────────────────────────────────
        ws2 = wb.create_sheet("Resumen")
        vigentes = sum(1 for c in certificados if c.estado_certificado == "vigente")
        por_vencer = sum(1 for c in certificados if c.estado_certificado == "por_vencer")
        vencidos = sum(1 for c in certificados if c.estado_certificado == "vencido")

        ws2["A1"] = "RESUMEN EJECUTIVO"
        ws2["A1"].font = Font(bold=True, size=14, color="1B4F8A")
        ws2["A3"] = "Total certificados:"
        ws2["B3"] = len(certificados)
        ws2["A4"] = "Vigentes:"
        ws2["B4"] = vigentes
        ws2["A5"] = "Por vencer (≤30 días):"
        ws2["B5"] = por_vencer
        ws2["A6"] = "Vencidos:"
        ws2["B6"] = vencidos
        ws2["A8"] = f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}"

        for row in range(3, 9):
            ws2[f"A{row}"].font = Font(bold=True)

        # Guardar en buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def exportar_excel_vencimientos(self, dias: int = 90) -> bytes:
        """Exporta solo certificados que vencen en los próximos N días."""
        from backend.models.certificado import Certificado
        hoy = date.today()
        certs = Certificado.query.filter(
            Certificado.fecha_vencimiento.between(hoy, hoy + timedelta(days=dias))
        ).order_by(Certificado.fecha_vencimiento.asc()).all()
        return self.exportar_excel_completo(certs)


report_service = ReportService()
